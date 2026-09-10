"""
Append one bot's results as a new column in XAUUSD_Bot_Comparison.xlsx, matching the
existing sheet's layout exactly (same row order, same formulas, same config/notes
sections). Always adds a NEW column to the right of the last bot - never inserts a
column in the middle, so every existing formula (which references fixed column
letters like B8, C8, ...) keeps working untouched.

Usage (edit the BOT dict below for each new run, or import add_bot() from your own
script and call it with the same fields):

    python add_bot_to_report.py

Requires: openpyxl (already used by the rest of this project's tooling).
After running, always call recalc.py (from the xlsx skill) on the file so the
formula cells get cached values again - Excel/LibreOffice will show blanks
otherwise until it's opened and recalculated once.
"""
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter

REPORT_PATH = r"C:\Users\lathu\Downloads\G\bot\reports\XAUUSD_Bot_Comparison.xlsx"
SHEET_NAME = "Bot Comparison"

# Fixed row numbers in the established layout - do not renumber unless the sheet's
# own structure changes (if it does, update these to match).
ROW_HEADER = 4          # 'Metric', bot names...
ROW_PNL = 5
ROW_PF = 6
ROW_TRADES = 7
ROW_WR = 8
ROW_LOSSRATE = 9        # formula, auto-generated
ROW_AVGWIN = 10
ROW_AVGLOSS = 11
ROW_EXPECTANCY = 12     # formula, auto-generated
ROW_BAL_DD_USD = 13
ROW_BAL_DD_PCT = 14
ROW_EQ_DD_USD = 15
ROW_EQ_DD_PCT = 16
ROW_DAYS = 17
ROW_PER_DAY = 18        # formula, auto-generated
ROW_TRADES_PER_DAY = 19 # formula, auto-generated

ROW_CONFIG_TITLE = 22
ROW_CONFIG_HEADER = 23
ROW_CONFIG_FIRST = 24   # first config-bullet row; extra rows inserted here if a bot
                        # needs more bullets than currently fit before "Notes:"


def _find_notes_row(ws):
    for row in range(ROW_CONFIG_FIRST, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "Notes:":
            return row
    raise RuntimeError("Could not find the 'Notes:' row - sheet layout may have changed")


def add_bot(name: str, pnl: float, pf: float, trades: int, win_rate: float,
            avg_win: float, avg_loss: float, bal_dd_usd: float, bal_dd_pct: float,
            eq_dd_usd: float, eq_dd_pct: float, backtest_days: int,
            config_bullets: list[str], extra_notes: list[str] | None = None,
            path: str = REPORT_PATH):
    """
    win_rate as a fraction (0.55, not 55). avg_loss NEGATIVE (e.g. -12.06).
    config_bullets: one string per row, in the order you want them to read top-to-bottom
    (they land under the bot's own column only - other bots' rows are untouched).
    extra_notes: optional list of new bullet lines appended to the end of Notes.
    """
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_NAME]

    # 1. find the next free column right after the current last bot column
    col = 2
    while ws.cell(row=ROW_HEADER, column=col).value not in (None, ""):
        col += 1
    letter = get_column_letter(col)

    # 2. metrics + formulas (formulas mirror the existing per-column pattern exactly)
    ws.cell(row=ROW_HEADER, column=col, value=name)
    ws.cell(row=ROW_PNL, column=col, value=pnl)
    ws.cell(row=ROW_PF, column=col, value=pf)
    ws.cell(row=ROW_TRADES, column=col, value=trades)
    ws.cell(row=ROW_WR, column=col, value=win_rate)
    ws.cell(row=ROW_LOSSRATE, column=col, value=f"=1-{letter}{ROW_WR}")
    ws.cell(row=ROW_AVGWIN, column=col, value=avg_win)
    ws.cell(row=ROW_AVGLOSS, column=col, value=avg_loss)
    ws.cell(row=ROW_EXPECTANCY, column=col,
            value=f"={letter}{ROW_WR}*{letter}{ROW_AVGWIN}+{letter}{ROW_LOSSRATE}*{letter}{ROW_AVGLOSS}")
    ws.cell(row=ROW_BAL_DD_USD, column=col, value=bal_dd_usd)
    ws.cell(row=ROW_BAL_DD_PCT, column=col, value=bal_dd_pct)
    ws.cell(row=ROW_EQ_DD_USD, column=col, value=eq_dd_usd)
    ws.cell(row=ROW_EQ_DD_PCT, column=col, value=eq_dd_pct)
    ws.cell(row=ROW_DAYS, column=col, value=backtest_days)
    ws.cell(row=ROW_PER_DAY, column=col, value=f"={letter}{ROW_PNL}/{letter}{ROW_DAYS}")
    ws.cell(row=ROW_TRADES_PER_DAY, column=col, value=f"={letter}{ROW_TRADES}/{letter}{ROW_DAYS}")

    # copy number formats from column B (the first bot column) so the new column
    # renders identically (currency/percent formatting etc.)
    for row in (ROW_PNL, ROW_PF, ROW_TRADES, ROW_WR, ROW_LOSSRATE, ROW_AVGWIN, ROW_AVGLOSS,
                ROW_EXPECTANCY, ROW_BAL_DD_USD, ROW_BAL_DD_PCT, ROW_EQ_DD_USD, ROW_EQ_DD_PCT,
                ROW_DAYS, ROW_PER_DAY, ROW_TRADES_PER_DAY):
        src = ws.cell(row=row, column=2)
        dst = ws.cell(row=row, column=col)
        dst.number_format = src.number_format
        dst.font = copy(src.font)

    # 3. config section - make sure there's room for every bullet before "Notes:"
    notes_row = _find_notes_row(ws)
    available_rows = notes_row - ROW_CONFIG_FIRST - 2  # leave the 2 blank spacer rows before Notes:
    if len(config_bullets) > available_rows:
        shortfall = len(config_bullets) - available_rows
        ws.insert_rows(notes_row, amount=shortfall)
        notes_row += shortfall

    ws.cell(row=ROW_CONFIG_HEADER, column=col, value=name)
    for i, bullet in enumerate(config_bullets):
        ws.cell(row=ROW_CONFIG_FIRST + i, column=col, value=bullet)

    # 4. optional extra notes, appended after the last existing note line
    if extra_notes:
        notes_row = _find_notes_row(ws)  # re-find in case rows shifted above
        last_note_row = notes_row
        r = notes_row + 1
        while ws.cell(row=r, column=1).value not in (None, ""):
            last_note_row = r
            r += 1
        for i, note in enumerate(extra_notes):
            ws.cell(row=last_note_row + 1 + i, column=1, value=note)

    wb.save(path)
    print(f"Added '{name}' as column {letter}. Now run recalc.py on the file.")


if __name__ == "__main__":
    # Example - edit these values for the bot you just validated, then run this file.
    add_bot(
        name="New Bot Name",
        pnl=0.0, pf=1.0, trades=0, win_rate=0.0,
        avg_win=0.0, avg_loss=0.0,
        bal_dd_usd=0.0, bal_dd_pct=0.0, eq_dd_usd=0.0, eq_dd_pct=0.0,
        backtest_days=982,
        config_bullets=[
            "Describe input 1: value",
            "Describe input 2: value",
        ],
        extra_notes=None,
    )
